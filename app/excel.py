from __future__ import annotations

import logging
import re
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models import Category, Creator, Video
from app.utils import channel_handle, normalize_channel_url, parse_sheet_category_id

log = logging.getLogger(__name__)

NAME_HEADERS = {"youtube channel", "channel name", "name"}
MODEL_HEADERS = {
    "monetization model",
    "modetization model",
    "plan",
    "model",
}
URL_HEADERS = {
    "channel links",
    "channel link",
    "channel url",
    "url",
    "link",
    "links",
    "channel",
    "column1",
}
CREATOR_ID_HEADERS = {"id", "creator id", "creator_id", "creatorid"}


def _header_map(values: list[str]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for index, raw in enumerate(values):
        key = (raw or "").strip().lower()
        if key in NAME_HEADERS:
            mapping["name"] = index
        elif key in MODEL_HEADERS:
            mapping["monetization_model"] = index
        elif key in URL_HEADERS:
            mapping["url"] = index
        elif key in CREATOR_ID_HEADERS:
            mapping["creator_id"] = index
    return mapping


def _cell_text(cell) -> str:
    value = cell.value
    if value is None:
        return ""
    return str(value).strip()


def _cell_int(cell) -> int | None:
    value = cell.value
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value)
    text = str(value).strip()
    digits = re.sub(r"[^\d]", "", text)
    return int(digits) if digits else None


def _looks_like_youtube_ref(value: str) -> bool:
    lowered = value.lower()
    return (
        "youtube.com/" in lowered
        or "youtu.be/" in lowered
        or value.startswith("@")
    )


def _cell_url(cell) -> str:
    text = _cell_text(cell)
    link = ""
    if cell.hyperlink and cell.hyperlink.target:
        link = str(cell.hyperlink.target).strip()
    if _looks_like_youtube_ref(text):
        return text
    if link:
        return link
    return text


def _clear_unuploaded_videos(session: Session, creator: Creator) -> int:
    videos = (
        session.query(Video)
        .filter(
            Video.creator_row_id == creator.id,
            Video.transfer_status.in_(
                {"pending", "failed", "skipped", "downloading", "uploading"}
            ),
        )
        .all()
    )
    for video in videos:
        session.delete(video)
    return len(videos)


def load_excel(session: Session, excel_path: Path) -> tuple[int, int]:
    if not excel_path.exists():
        raise FileNotFoundError(
            f"Excel file not found: {excel_path}. Put the workbook at this path or set EXCEL_PATH."
        )

    workbook = load_workbook(excel_path, data_only=True, read_only=False)
    categories_upserted = 0
    creators_upserted = 0

    for sheet in workbook.worksheets:
        category_name = (sheet.title or "").strip()
        if not category_name:
            continue

        sheet_category_id = parse_sheet_category_id(category_name)
        if sheet_category_id is None:
            log.warning(
                "Skipping sheet '%s' — category id missing. "
                "Name the sheet like 'Kids - MiniMinds (1)'.",
                category_name,
            )
            continue

        rows = list(sheet.iter_rows())
        if not rows:
            log.warning("Skipping empty sheet: %s", category_name)
            continue

        headers = [_cell_text(cell) for cell in rows[0]]
        columns = _header_map(headers)
        if "name" not in columns or "url" not in columns or "creator_id" not in columns:
            log.warning(
                "Sheet '%s' needs columns: YouTube Channel, Channel links, Id. Found: %s",
                category_name,
                headers,
            )
            continue

        category = session.get(Category, sheet_category_id)
        if category is None:
            category = Category(id=sheet_category_id, name=category_name)
            session.add(category)
            session.flush()
        else:
            category.name = category_name
        categories_upserted += 1

        for row in rows[1:]:
            name = _cell_text(row[columns["name"]])
            url = normalize_channel_url(_cell_url(row[columns["url"]]))
            creator_id = _cell_int(row[columns["creator_id"]])
            model = None
            if "monetization_model" in columns:
                model = _cell_text(row[columns["monetization_model"]]) or None
            if not name or not url:
                continue
            if creator_id is None:
                log.warning(
                    "Skipping '%s' on sheet '%s' — missing creator Id",
                    name,
                    category_name,
                )
                continue

            handle = channel_handle(url)
            creator = (
                session.query(Creator)
                .filter_by(category_id=category.id, creator_id=creator_id)
                .one_or_none()
            )
            if creator is None:
                creator = (
                    session.query(Creator)
                    .filter_by(category_id=category.id, channel_url=url)
                    .one_or_none()
                )

            if creator is None:
                creator = Creator(
                    category_id=category.id,
                    creator_id=creator_id,
                    name=name,
                    monetization_model=model,
                    channel_url=url,
                    handle=handle,
                    scrape_status="pending",
                )
                session.add(creator)
                creators_upserted += 1
            else:
                if creator.channel_url != url:
                    log.warning(
                        "Updating URL for creator_id=%s '%s': %s -> %s",
                        creator_id,
                        name,
                        creator.channel_url,
                        url,
                    )
                    removed = _clear_unuploaded_videos(session, creator)
                    if removed:
                        log.info(
                            "Removed %s old videos for '%s' after URL fix",
                            removed,
                            name,
                        )
                    creator.channel_url = url
                    creator.handle = handle
                    creator.youtube_channel_id = None
                    creator.bunny_collection_id = None
                    creator.scrape_status = "pending"
                    creator.scrape_error = None
                creator.creator_id = creator_id
                creator.name = name
                creator.monetization_model = model
                if handle:
                    creator.handle = handle

        log.info(
            "Loaded sheet '%s' (category_id=%s)",
            category_name,
            sheet_category_id,
        )

    session.flush()
    return categories_upserted, creators_upserted
